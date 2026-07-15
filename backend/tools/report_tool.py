"""tools/report_tool.py — generate Word (.docx) and Excel (.xlsx) reports"""

import os, logging
from datetime import datetime
from langchain_core.tools import Tool
from langchain_core.tools import StructuredTool
from pydantic import BaseModel
log = logging.getLogger(__name__)
OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

class WordReportInput(BaseModel):
    content: str
    filename: str = ""
    title: str = "Agent Report"
class ExcelReportInput(BaseModel):
    data: str
    filename: str = ""
    sheet_name: str = "Report"
def _out_path(filename, ext):
    if not filename:
        filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
    filename = os.path.basename(filename)  # strip any directory components
    if not filename.endswith(ext):
        filename += ext
    return os.path.abspath(os.path.join(OUTPUT_DIR, filename))


def save_word_report(content: str, filename: str = "", title: str = "Agent Report") -> str:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        filename = _out_path(filename, ".docx")
        doc = Document()
        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        ts = doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y at %H:%M')}")
        ts.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ts.runs[0].font.size = Pt(10)
        ts.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        doc.add_paragraph()

        for line in content.split("\n"):
            s = line.strip()
            if not s:
                doc.add_paragraph()
            elif s.startswith("## "):
                doc.add_heading(s[3:], level=2)
            elif s.startswith("# "):
                doc.add_heading(s[2:], level=1)
            elif s.startswith(("- ", "* ")):
                doc.add_paragraph(s[2:], style="List Bullet")
            else:
                doc.add_paragraph(s)

        doc.save(filename)
        return f"Word report saved to {filename}"
    except ImportError:
        return "Error: python-docx not installed."
    except Exception as e:
        return f"Error creating Word report: {e}"


def save_excel_report(data: str, filename: str = "", sheet_name: str = "Report") -> str:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        filename = _out_path(filename, ".xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        lines = [l.strip() for l in data.strip().split("\n") if l.strip()]
        rows = []
        for line in lines:
            if "|" in line:
                cells = [c.strip() for c in line.split("|") if c.strip()]
            else:
                import csv, io
                cells = next(csv.reader(io.StringIO(line)), [line])
            rows.append(cells)

        if not rows:
            return "No data to write."

        hfill = PatternFill("solid", fgColor="1F497D")
        hfont = Font(bold=True, color="FFFFFF", size=11)
        border = Border(bottom=Side(style="thin", color="CCCCCC"))
        alt = PatternFill("solid", fgColor="EEF2F7")

        for ci, v in enumerate(rows[0], 1):
            c = ws.cell(row=1, column=ci, value=v)
            c.fill = hfill; c.font = hfont
            c.alignment = Alignment(horizontal="center", vertical="center")

        for ri, row in enumerate(rows[1:], 2):
            for ci, v in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=v)
                c.border = border
                if ri % 2 == 0:
                    c.fill = alt
                c.alignment = Alignment(vertical="center")

        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(w + 4, 50)

        ws.row_dimensions[1].height = 22
        wb.save(filename)
        return f"Excel report saved to {filename} ({len(rows)-1} rows)"
    except ImportError:
        return "Error: openpyxl not installed."
    except Exception as e:
        return f"Error creating Excel report: {e}"


save_word_tool = StructuredTool.from_function(
    func=save_word_report,
    name="save_word_report",
    description="Save content as a Word (.docx) document.",
    args_schema=WordReportInput,
)

save_excel_tool = StructuredTool.from_function(
    func=save_excel_report,
    name="save_excel_report",
    description="Save tabular data as Excel (.xlsx).",
    args_schema=ExcelReportInput,
)
